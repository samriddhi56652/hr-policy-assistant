"""processor — turn raw files (txt/pdf/docx/pptx) into plain text, once.

    raw/<prefix>/file.pdf  --parse-->  processed/<prefix>/file.json
"""

import io
import json

from docx import Document as DocxReader
from google.cloud import storage
from pptx import Presentation
from pypdf import PdfReader

from hr_assistant import config

# Which raw folder feeds which processed folder.
_RAW_TO_PROCESSED = {
    config.GCS_PREFIX: config.PROCESSED_HR_PREFIX,
    config.NOISE_GCS_PREFIX: config.PROCESSED_NOISE_PREFIX,
}


def extract_policy_category(text: str) -> str:
    """HR files have a 'Policy Category: X' line, noise files a 'Category: X'
    line, near the top. Return X, or 'Unknown' if neither is found."""
    for line in text.splitlines()[:5]:
        stripped = line.strip().lower()
        if stripped.startswith("policy category:") or stripped.startswith("category:"):
            return line.split(":", 1)[1].strip()
    return "Unknown"


def _parse_pdf(raw_bytes: bytes) -> str:
    reader = PdfReader(io.BytesIO(raw_bytes))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def _parse_docx(raw_bytes: bytes) -> str:
    doc = DocxReader(io.BytesIO(raw_bytes))
    return "\n".join(p.text for p in doc.paragraphs)


def _parse_pptx(raw_bytes: bytes) -> str:
    prs = Presentation(io.BytesIO(raw_bytes))
    lines = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                lines.append(shape.text_frame.text)
    return "\n".join(lines)


_PARSERS = {".pdf": _parse_pdf, ".docx": _parse_docx, ".pptx": _parse_pptx}


def parse_blob(blob) -> str:
    """Download one file from the bucket and return its plain text."""
    ext = "." + blob.name.rsplit(".", 1)[-1].lower()
    if ext == ".txt":
        return blob.download_as_text()
    parser = _PARSERS.get(ext)
    if parser is None:
        raise ValueError(f"No parser for {ext!r} ({blob.name})")
    return parser(blob.download_as_bytes())


def process_raw_to_json() -> int:
    """Parse every raw file and save it as a JSON record in processed/.
    Returns how many records were written."""
    bucket = storage.Client(project=config.PROJECT_ID).bucket(config.GCS_BUCKET_NAME)

    count = 0
    for raw_prefix, processed_prefix in _RAW_TO_PROCESSED.items():
        for blob in bucket.list_blobs(prefix=raw_prefix):
            filename = blob.name.rsplit("/", 1)[-1]
            if "." not in filename:
                continue  # skip folder placeholders
            text = parse_blob(blob)
            text = text.replace("\r\n", "\n")  # Windows line endings -> plain \n, so paragraph splitting works
            record = {
                "source": filename,
                "policy_category": extract_policy_category(text),
                "text": text,
                "raw_gcs_path": f"gs://{config.GCS_BUCKET_NAME}/{blob.name}",
            }
            json_name = filename.rsplit(".", 1)[0] + ".json"
            bucket.blob(processed_prefix + json_name).upload_from_string(
                json.dumps(record, indent=2), content_type="application/json"
            )
            count += 1
    return count